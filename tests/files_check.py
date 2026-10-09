"""The file pages, end to end, with an isolated relay and agent and a made-up project: what a folder lists, that a
file comes back whole, and that nothing outside the project can be reached. With a browser it also takes a picture
of the folder and of every kind of file. Build the agent first.

    python tests/files_check.py [chrome.exe <folder for the pictures>]

The Word, Excel, PowerPoint and PDF samples are made with python-docx, openpyxl, python-pptx and reportlab when
they are installed; without them those kinds are left out.
"""
import base64
import hashlib
import http.client
import json
import os
import secrets
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from urllib.parse import quote

sys.path.insert(0, str(Path(__file__).resolve().parent))
from ws_client import Socket

ROOT = Path(__file__).resolve().parents[1]
PORT = int(os.environ.get("RCLI_TEST_PORT", "8737"))
password = "test-" + secrets.token_hex(8)

PYTHON = '''"""Relative risk with a confidence interval."""
import math


def relative_risk(a, b, c, d, level=0.95):
    # a, b: exposed with and without the outcome; c, d: unexposed
    risk = (a / (a + b)) / (c / (c + d))
    se = math.sqrt(1 / a - 1 / (a + b) + 1 / c - 1 / (c + d))
    z = 1.959964 if level == 0.95 else 2.575829
    return risk, risk * math.exp(-z * se), risk * math.exp(z * se)


print("RR = %.2f (%.2f, %.2f)" % relative_risk(30, 70, 15, 85))
'''
MARKDOWN = '''---
title: 分析说明
date: 2026-10-09
---

# 队列研究分析说明

本项目比较 **暴露组** 与 *非暴露组* 的结局发生率，脚本见 [analysis.py](analysis.py)。

## 主要结果

| 分组 | 人数 | 事件数 | 发生率 |
| --- | ---: | ---: | ---: |
| 暴露组 | 100 | 30 | 30.0% |
| 非暴露组 | 100 | 15 | 15.0% |

> 相对危险度 RR = 2.00（95% CI 1.15–3.48）。

```r
fit <- glm(event ~ exposed + age, family = binomial, data = cohort)
summary(fit)
```

- [x] 数据清洗
- [ ] 敏感性分析

![趋势图](figures/trend.png)
'''
NOTEBOOK = {"cells": [
    {"cell_type": "markdown", "metadata": {}, "source": ["## 描述统计\n", "先看年龄的分布。"]},
    {"cell_type": "code", "execution_count": 1, "metadata": {}, "source": ["import statistics\n", "ages = [34, 41, 29, 52, 47]\n", "statistics.mean(ages)"],
     "outputs": [{"output_type": "execute_result", "data": {"text/plain": ["40.6"]}, "execution_count": 1, "metadata": {}}]}],
    "metadata": {"language_info": {"name": "python"}}, "nbformat": 4, "nbformat_minor": 5}


def picture(path):
    from PIL import Image, ImageDraw
    image = Image.new("RGB", (640, 360), "#f4f6fb")
    draw = ImageDraw.Draw(image)
    points = [(40 + n * 56, 300 - int(180 * (0.2 + 0.07 * n + 0.05 * ((n * 7) % 3)))) for n in range(11)]
    draw.line([(40, 310), (610, 310)], fill="#9aa3b5", width=2)
    draw.line(points, fill="#2f6fe4", width=5)
    for x, y in points:
        draw.ellipse([x - 6, y - 6, x + 6, y + 6], fill="#2f6fe4")
    image.save(path)


def samples(project):
    """Fills the project; returns the files that were made, by the kind of view they should get."""
    made = {}
    (project / "figures").mkdir()
    (project / "data").mkdir()
    (project / ".git").mkdir()
    (project / ".git" / "HEAD").write_text("ref: refs/heads/main\n")
    (project / "analysis.py").write_text(PYTHON, encoding="utf-8")
    (project / "README.md").write_text(MARKDOWN, encoding="utf-8")
    (project / "explore.ipynb").write_text(json.dumps(NOTEBOOK), encoding="utf-8")
    (project / "data" / "cohort.csv").write_bytes("编号,分组,年龄,结局\n1,暴露组,34,1\n2,非暴露组,41,0\n3,暴露组,52,1\n".encode("gbk"))
    (project / "model.bin").write_bytes(bytes(range(256)) * 8)
    (project / "notes").write_text("没有后缀的文本文件\n第二行\n", encoding="utf-8")
    (project / "big.log").write_bytes(b"0123456789abcdef" * 150_000)          # 2.4 MB: more than one piece
    (project / "bundle.zip").write_bytes(os.urandom(1_600_000))               # not shown, only saved: three pieces
    made.update(code="analysis.py", markdown="README.md", notebook="explore.ipynb", sheet_csv="data/cohort.csv", unknown="model.bin")
    try:
        picture(project / "figures" / "trend.png")
        made["image"] = "figures/trend.png"
    except ImportError:
        pass
    try:
        import docx
        document = docx.Document()
        document.add_heading("研究方案摘要", 1)
        document.add_paragraph("本研究为回顾性队列研究，纳入 2020 年至 2024 年就诊的成年患者，主要结局为 30 天内再入院。")
        document.add_heading("统计分析", 2)
        document.add_paragraph("采用 Cox 比例风险模型估计风险比及其 95% 置信区间。", style="List Bullet")
        table = document.add_table(rows=3, cols=3)
        table.style = "Table Grid"
        for r, row in enumerate([("变量", "暴露组", "非暴露组"), ("年龄（岁）", "52.3", "51.8"), ("女性（%）", "46.0", "44.5")]):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        document.save(project / "方案.docx")
        made["word"] = "方案.docx"
    except ImportError:
        pass
    try:
        import openpyxl
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "基线"
        for row in [("变量", "暴露组", "非暴露组", "P 值"), ("年龄", 52.3, 51.8, 0.412), ("BMI", 24.6, 24.1, 0.087), ("吸烟", 0.31, 0.27, 0.203)]:
            sheet.append(row)
        other = book.create_sheet("结局")
        other.append(("分组", "事件数", "人数"))
        other.append(("暴露组", 30, 100))
        book.save(project / "data" / "结果.xlsx")
        made["sheet"] = "data/结果.xlsx"
    except ImportError:
        pass
    try:
        import pptx
        deck = pptx.Presentation()
        first = deck.slides.add_slide(deck.slide_layouts[0])
        first.shapes.title.text = "组会汇报"
        first.placeholders[1].text = "队列研究的初步结果"
        second = deck.slides.add_slide(deck.slide_layouts[1])
        second.shapes.title.text = "主要发现"
        second.placeholders[1].text = "暴露组发生率 30.0%\n非暴露组发生率 15.0%\n相对危险度 2.00"
        if "image" in made:
            second.shapes.add_picture(str(project / "figures" / "trend.png"), pptx.util.Inches(5), pptx.util.Inches(3), width=pptx.util.Inches(4))
        deck.save(project / "汇报.pptx")
        made["slides"] = "汇报.pptx"
    except ImportError:
        pass
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        page = canvas.Canvas(str(project / "paper.pdf"), pagesize=A4)
        for n in range(3):
            page.setFont("Helvetica-Bold", 20)
            page.drawString(72, 760, "Cohort analysis, page %d" % (n + 1))
            page.setFont("Helvetica", 12)
            for line in range(18):
                page.drawString(72, 720 - line * 22, "Line %d of the report: the relative risk was 2.00 (95%% CI 1.15 to 3.48)." % (line + 1))
            page.showPage()
        page.save()
        made["pdf"] = "paper.pdf"
    except ImportError:
        pass
    return made


def main():
    chrome, out = (sys.argv[1], Path(sys.argv[2])) if len(sys.argv) > 2 else (None, None)
    temp = tempfile.TemporaryDirectory()
    data, profile, project, outside = (Path(temp.name) / name for name in ("agent", "profile", "cohort-study", "outside"))
    for folder in (data, project, outside):
        folder.mkdir()
    made = samples(project)
    # RCLI_TEST_DOCX and RCLI_TEST_PPTX put a document of one's own in place of the made-up ones, to see how it is shown.
    for wanted, kind in (("RCLI_TEST_DOCX", "word"), ("RCLI_TEST_PPTX", "slides")):
        if os.environ.get(wanted) and kind in made:
            (project / made[kind]).write_bytes(Path(os.environ[wanted]).read_bytes())
    (outside / "secret.txt").write_text("not part of the project")
    linked = subprocess.run(["cmd", "/c", "mklink", "/J", str(project / "link"), str(outside)], capture_output=True).returncode == 0
    (data / "config.json").write_text(json.dumps({"Server": "http://127.0.0.1:%d" % PORT, "RemoteEnabled": True, "RemoteMaxMode": "full", "RemoteDirs": ["cohort-study=%s" % project]}), encoding="utf-8")
    hidden = {"creationflags": subprocess.CREATE_NO_WINDOW}
    relay = subprocess.Popen([sys.executable, str(ROOT / "relay" / "server.py"), "--port", str(PORT), "--data", str(Path(temp.name) / "relay")],
                             env=dict(os.environ, RCLI_PASSWORD=password, PYTHONUTF8="1"), stdout=subprocess.DEVNULL, **hidden)
    agent_exe = ROOT / "agent-windows" / "bin" / "RemoteCliAgent.exe"
    subprocess.run([str(agent_exe), "--set-password", str(data)], input=password.encode(), check=True, **hidden)
    agent = subprocess.Popen([str(agent_exe), "--data", str(data)], **hidden)
    try:
        cookie = {}

        def call(path, payload=None, expect=200):
            connection = http.client.HTTPConnection("127.0.0.1", PORT, timeout=40)
            body = None if payload is None else json.dumps(payload).encode()
            connection.request("POST" if body is not None else "GET", path, body=body, headers=dict(cookie, **({"Content-Type": "application/json"} if body is not None else {})))
            reply = connection.getresponse()
            text = reply.read()
            if reply.getheader("Set-Cookie"):
                cookie["Cookie"] = reply.getheader("Set-Cookie").split(";")[0]
            assert reply.status == expect, (path, reply.status, text[:200].decode("utf-8", "replace"))
            return json.loads(text)

        files = lambda action, path="", offset=0, expect=200, folder="cohort-study": call("/api/files", {"id": secrets.token_hex(16), "action": action, "dir": folder, "path": path, "offset": offset}, expect)
        for _ in range(40):
            try:
                call("/api/login", {"password": password})
                break
            except OSError:
                time.sleep(0.5)
        for _ in range(60):
            if call("/api/terminal")["device"]["workspaces"]:
                break
            time.sleep(0.5)
        features = call("/api/terminal")["device"]["features"]
        assert "files" in features

        top = files("file_list")
        names = {e["name"]: e for e in top["entries"]}
        assert names["figures"]["dir"] and not names["analysis.py"]["dir"] and names["analysis.py"]["size"] == (project / "analysis.py").stat().st_size, names
        assert names[".git"]["hidden"] and not names["README.md"]["hidden"] and abs(names["README.md"]["modified"] / 1000 - time.time()) < 600
        assert [e["name"] for e in files("file_list", "data")["entries"]][0] == "cohort.csv"
        # a file larger than one piece arrives whole
        whole, at = b"", 0
        while True:
            piece = files("file_read", "big.log", at)
            whole += base64.b64decode(piece["data"])
            at = len(whole)
            if piece["end"]:
                break
        assert piece["size"] == 2_400_000 and hashlib.sha256(whole).digest() == hashlib.sha256((project / "big.log").read_bytes()).digest()
        assert base64.b64decode(files("file_read", "data/cohort.csv")["data"]).decode("gbk").startswith("编号")
        # nothing outside the project
        for path in ("..", "../outside/secret.txt", "data/../../outside/secret.txt", "C:\\Windows\\win.ini", "\\\\localhost\\c$\\Windows\\win.ini"):
            for action in ("file_list", "file_read"):
                error = files(action, path, expect=400)["error"]
                assert "项目文件夹" in error or "无效" in error or "不存在" in error, (path, error)
        if linked:
            assert "之外" in files("file_list", "link", expect=400)["error"] and "之外" in files("file_read", "link/secret.txt", expect=400)["error"]
        assert "没有这个项目" in files("file_list", folder="other", expect=400)["error"]
        assert "不存在" in files("file_read", "missing.txt", expect=400)["error"]
        assert "文件夹" in files("file_read", "data", expect=400)["error"] and "不是文件夹" in files("file_list", "analysis.py", expect=400)["error"]
        cookie.clear()
        call("/api/files", {"id": secrets.token_hex(16), "action": "file_list", "dir": "cohort-study", "path": ""}, 401)
        report = {"listed": len(top["entries"]), "pieces_of_big_file": -(-2_400_000 // 737_280), "outside_refused": True, "junction_refused": linked, "samples": sorted(made)}

        if chrome:
            out.mkdir(parents=True, exist_ok=True)
            # The browser is driven in real time: a PDF is drawn by a helper thread that a browser on "virtual time" never runs.
            browser = subprocess.Popen([chrome, "--headless=new", "--enable-unsafe-swiftshader", "--hide-scrollbars", "--no-first-run", "--user-data-dir=" + str(profile),
                                        "--remote-debugging-port=%d" % (PORT + 1), "--remote-allow-origins=*", "--window-size=500,960", "about:blank"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            try:
                for _ in range(60):
                    try:
                        connection = http.client.HTTPConnection("127.0.0.1", PORT + 1, timeout=5)
                        connection.request("PUT", "/json/new?about:blank")
                        page = json.loads(connection.getresponse().read())
                        break
                    except (OSError, ValueError):
                        time.sleep(0.5)
                pipe = Socket("http://127.0.0.1:%d" % (PORT + 1), "/devtools/page/" + page["id"], "", 20)
                sent, said = [0], []

                def ask(method, **params):
                    sent[0] += 1
                    pipe.send({"id": sent[0], "method": method, "params": params})
                    while True:
                        message = pipe.receive(30)
                        if message.get("method") == "Runtime.consoleAPICalled" and message["params"]["type"] in ("error", "warning"):
                            said.append(" ".join(str(a.get("value", a.get("description", ""))) for a in message["params"]["args"])[:300])
                        if message.get("method") == "Runtime.exceptionThrown":
                            said.append(str(message["params"]["exceptionDetails"].get("exception", {}).get("description", message["params"]["exceptionDetails"].get("text")))[:300])
                        if message.get("id") == sent[0]:
                            return message.get("result", {})

                ask("Page.enable")
                ask("Runtime.enable")
                ask("Emulation.setDeviceMetricsOverride", width=500, height=960, deviceScaleFactor=1, mobile=True)

                def shot(url, name, ready="true", seconds=20):
                    """Opens the page, waits until `ready` (an expression in the page) holds and nothing is loading, and saves a picture."""
                    del said[:]
                    ask("Page.navigate", url=url)
                    deadline = time.time() + seconds
                    settled = "document.readyState === 'complete' && !!document.getElementById('loading') && document.getElementById('loading').hidden !== false && (%s)" % ready
                    while time.time() < deadline:
                        time.sleep(0.4)
                        if ask("Runtime.evaluate", expression=settled, returnByValue=True).get("result", {}).get("value") is True:
                            break
                    else:
                        raise AssertionError("%s did not finish: %s" % (name, said[:3]))
                    time.sleep(0.6)
                    (out / name).write_bytes(base64.b64decode(ask("Page.captureScreenshot", format="png")["data"]))
                    assert not said, (name, said[:3])        # the page reported a problem of its own

                skin = os.environ.get("RCLI_TEST_SKIN", "")
                base = "http://127.0.0.1:%d" % PORT
                ask("Page.navigate", url=base + "/" + ("?skin=" + skin if skin else "") + "#p=" + password)       # signs in and keeps the skin
                time.sleep(3)
                view = lambda path: base + "/files/?dir=cohort-study" + ("&path=" + quote(path.rsplit("/", 1)[0]) if "/" in path else "") + "&file=" + quote(path.rsplit("/", 1)[-1])
                drawn = {"pdf": "document.querySelectorAll('.sheet-of-paper canvas').length > 0", "image": "!!document.querySelector('.picture img') && document.querySelector('.picture img').naturalWidth > 0",
                         "markdown": "Array.from(document.querySelectorAll('.doc img')).every(i => i.naturalWidth > 0) && !!document.querySelector('.doc h1')",
                         "word": "document.querySelectorAll('.word-looks section.docx table').length > 0", "sheet": "document.querySelectorAll('.grid td').length > 8", "sheet_csv": "document.querySelectorAll('.grid td').length > 8",
                         "slides": "document.querySelectorAll('.slide-looks .pptx-preview-wrapper > *').length >= 2", "notebook": "!!document.querySelector('.notebook .code')", "code": "!!document.querySelector('.code .hljs-keyword')",
                         "unknown": "!!document.querySelector('.unknown')"}
                shot(base + "/?project=cohort-study", "project.png", "!document.getElementById('project').hidden")
                shot(base + "/files/?dir=cohort-study", "folder.png", "document.querySelectorAll('#entries li').length > 5")
                for kind, path in sorted(made.items()):
                    shot(view(path), kind + ".png", drawn[kind], seconds=60 if kind in ("word", "slides") else 20)
                # Saving inside the app: the page hands the file to the phone piece by piece. A stand-in for the app
                # collects the pieces; together they must be the file.
                stand_in = ask("Page.addScriptToEvaluateOnNewDocument", source="window.RemoteCliNative = { got: [], saveStart(n) { this.name = n; this.got = []; return ''; }, "
                               "savePiece(d) { this.got.push(d); return ''; }, saveEnd() { this.done = true; return '下载 / RemoteCLI'; }, saveCancel() { this.cancelled = true; }, openSaved() {}, chrome() {}, pref() { return ''; }, setPref() {} };")
                shot(view("bundle.zip"), "save.png", "!!document.querySelector('.unknown button')")
                ask("Runtime.evaluate", expression="document.querySelector('.unknown button').click()")
                deadline = time.time() + 30
                while time.time() < deadline and ask("Runtime.evaluate", expression="window.RemoteCliNative.done === true", returnByValue=True).get("result", {}).get("value") is not True:
                    time.sleep(0.3)
                digest = ask("Runtime.evaluate", awaitPromise=True, returnByValue=True, expression="(async () => { const n = window.RemoteCliNative, raw = n.got.map(p => atob(p)).join(''), bytes = Uint8Array.from(raw, c => c.charCodeAt(0)); "
                             "const hash = await crypto.subtle.digest('SHA-256', bytes); return n.name + ' ' + n.got.length + ' ' + Array.from(new Uint8Array(hash), b => b.toString(16).padStart(2, '0')).join(''); })()")["result"]["value"]
                assert digest == "bundle.zip 3 " + hashlib.sha256((project / "bundle.zip").read_bytes()).hexdigest(), digest
                time.sleep(0.5)
                (out / "save.png").write_bytes(base64.b64decode(ask("Page.captureScreenshot", format="png")["data"]))
                ask("Page.removeScriptToEvaluateOnNewDocument", identifier=stand_in["identifier"])
                report["saved_in_pieces"] = 3
                pipe.close()
            finally:
                browser.kill()
            report["pictures"] = sorted(p.name for p in out.glob("*.png"))
        print(json.dumps(report, ensure_ascii=False))
    finally:
        agent.kill()
        relay.kill()
        time.sleep(1)
        subprocess.run(["cmd", "/c", "rmdir", str(project / "link")], capture_output=True)       # the link, not what it leads to
        try:
            temp.cleanup()
        except OSError:
            pass


if __name__ == "__main__":
    main()
